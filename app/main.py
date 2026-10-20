from __future__ import annotations

import io, json, secrets, hashlib, hmac, re, os
from datetime import datetime
from pathlib import Path
from typing import Any
from xml.sax.saxutils import escape

from fastapi import Depends, FastAPI, File, Header, HTTPException, Query, UploadFile
from fastapi.responses import FileResponse, HTMLResponse, StreamingResponse
from fastapi.staticfiles import StaticFiles
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from pydantic import BaseModel, Field
from app.classifier import classify_text

ROOT = Path(__file__).resolve().parent
DATA = Path(os.environ.get("HOSPITAL_DATA_DIR", ROOT / "data")).expanduser()
DATA.mkdir(exist_ok=True)
DB = DATA / "records.json"
USERS_DB = DATA / "users.json"
ADMIN_USERNAME = os.environ.get("ADMIN_USERNAME", "admin").strip().lower() or "admin"
ADMIN_PASSWORD = os.environ.get("ADMIN_PASSWORD", "demo1234")
DIAGNOSIS_GROUPS = {
    "01": "Brain", "02": "Psychosocial disorder", "03": "Spinal cord",
    "04": "Roots and nerves", "05": "Neuromuscular junction", "06": "Muscle disease",
    "07": "Congenital myopathies", "08": "Metabolic myopathies", "09": "Chromosome & syndromes",
    "10": "Endocrine & metabolism", "11": "Mitochondrial", "12": "Neurocutaneous",
    "13": "Vision", "14": "Auditory", "15": "Unspecified delay", "99": "Others",
}
app = FastAPI(title="Hospital Checklist Prototype")
app.mount("/static", StaticFiles(directory=ROOT / "static"), name="static")
TOKENS: dict[str, str] = {}


class Login(BaseModel):
    username: str
    password: str


class NewUser(BaseModel):
    username: str
    password: str
    role: str = "user"


class PasswordReset(BaseModel):
    password: str


class ClassifyRequest(BaseModel):
    text: str = Field(max_length=50000)


class Record(BaseModel):
    id: str | None = None
    caseId: str = ""
    patientName: str = ""
    chartNo: str = ""
    examDate: str = ""
    physician: str = ""
    functional: list[str] = Field(default_factory=list)
    diagnoses: list[str] = Field(default_factory=list)
    diagnosisLabels: list[str] = Field(default_factory=list)
    classificationReview: list[dict[str, str]] = Field(default_factory=list)
    classificationStatus: str = "review"
    diagnosisNotes: str = ""
    management: list[str] = Field(default_factory=list)
    medication: str = ""
    sourceText: str = ""


def _password_hash(password: str, salt: bytes | None = None) -> str:
    salt = salt or secrets.token_bytes(16)
    digest = hashlib.pbkdf2_hmac("sha256", password.encode(), salt, 310_000)
    return f"pbkdf2_sha256${salt.hex()}${digest.hex()}"


def _verify_password(password: str, encoded: str) -> bool:
    try:
        algorithm, salt_hex, digest_hex = encoded.split("$")
        if algorithm != "pbkdf2_sha256":
            return False
        candidate = hashlib.pbkdf2_hmac("sha256", password.encode(), bytes.fromhex(salt_hex), 310_000).hex()
        return hmac.compare_digest(candidate, digest_hex)
    except (ValueError, TypeError):
        return False


def _database_url() -> str | None:
    return os.environ.get("DATABASE_URL")


def _pg_get(key: str):
    import psycopg
    url = _database_url()
    with psycopg.connect(url.replace("postgres://", "postgresql://", 1)) as conn:
        conn.execute("CREATE TABLE IF NOT EXISTS app_state (key text PRIMARY KEY, value jsonb NOT NULL)")
        row = conn.execute("SELECT value FROM app_state WHERE key = %s", (key,)).fetchone()
        return row[0] if row else None


def _pg_set(key: str, value: Any):
    import psycopg
    from psycopg.types.json import Jsonb
    url = _database_url()
    with psycopg.connect(url.replace("postgres://", "postgresql://", 1)) as conn:
        conn.execute("CREATE TABLE IF NOT EXISTS app_state (key text PRIMARY KEY, value jsonb NOT NULL)")
        conn.execute(
            "INSERT INTO app_state (key, value) VALUES (%s, %s) "
            "ON CONFLICT (key) DO UPDATE SET value = EXCLUDED.value",
            (key, Jsonb(value)),
        )


def load_users() -> list[dict[str, str]]:
    if _database_url():
        users = _pg_get("users")
        if users is not None:
            return users
        initial = [{"username": ADMIN_USERNAME, "role": "admin", "passwordHash": _password_hash(ADMIN_PASSWORD)}]
        _pg_set("users", initial)
        return initial
    if USERS_DB.exists():
        return json.loads(USERS_DB.read_text())
    initial = [{"username": ADMIN_USERNAME, "role": "admin", "passwordHash": _password_hash(ADMIN_PASSWORD)}]
    USERS_DB.write_text(json.dumps(initial, indent=2))
    return initial


def save_users(users: list[dict[str, str]]):
    if _database_url():
        _pg_set("users", users)
        return
    USERS_DB.write_text(json.dumps(users, ensure_ascii=False, indent=2))


def auth(authorization: str | None = Header(default=None)) -> dict[str, str]:
    token = authorization.removeprefix("Bearer ") if authorization else ""
    username = TOKENS.get(token)
    user = next((u for u in load_users() if u["username"] == username), None)
    if user is None:
        raise HTTPException(401, "Sign in required")
    return user


def admin_auth(user: dict[str, str] = Depends(auth)) -> dict[str, str]:
    if user["role"] != "admin":
        raise HTTPException(403, "Administrator access required")
    return user


def read_records() -> list[dict[str, Any]]:
    if _database_url():
        return _pg_get("records") or []
    return json.loads(DB.read_text()) if DB.exists() else []


def write_records(rows: list[dict[str, Any]]):
    if _database_url():
        _pg_set("records", rows)
        return
    DB.write_text(json.dumps(rows, ensure_ascii=False, indent=2))


def extract_basic_fields(text: str) -> dict[str, str]:
    """Conservative label-based extraction only; no clinical diagnosis inference."""
    patterns = {
        "patientName": r"(?:patient\s*name|姓名|病患姓名)\s*[:：=|]\s*([^\r\n|;,]+)",
        "chartNo": r"(?:chart\s*(?:no\.?|number)|medical\s*record\s*(?:no\.?|number)|病歷號碼?|病歷編號)\s*[:：=|]\s*([^\r\n|;,]+)",
        "examDate": r"(?:examination\s*date|exam\s*date|檢查日期|評估日期)\s*[:：=|]\s*([^\r\n|;,]+)",
        "physician": r"(?:physician|doctor|醫師|醫師簽名)\s*[:：=|]\s*([^\r\n|;,]+)",
    }
    result: dict[str, str] = {}
    for key, pattern in patterns.items():
        match = re.search(pattern, text, re.IGNORECASE)
        if not match:
            continue
        value = match.group(1).strip()
        if key == "examDate":
            normalized = re.search(r"\b(\d{4})\s*[/.-]\s*(\d{1,2})\s*[/.-]\s*(\d{1,2})\b", value)
            if normalized:
                try:
                    value = datetime(int(normalized[1]), int(normalized[2]), int(normalized[3])).strftime("%Y-%m-%d")
                except ValueError:
                    continue
            else:
                continue
        result[key] = value
    return result


@app.get("/", response_class=HTMLResponse)
def home():
    return FileResponse(ROOT / "static" / "index.html")


@app.get("/style.css")
def stylesheet():
    return FileResponse(ROOT / "static" / "style.css", media_type="text/css")


@app.get("/app.js")
def frontend_script():
    return FileResponse(ROOT / "static" / "app.js", media_type="text/javascript")


@app.get("/diagnosis-data.js")
def diagnosis_data_script():
    return FileResponse(ROOT / "static" / "diagnosis-data.js", media_type="text/javascript")


@app.post("/api/login")
def login(body: Login):
    username = body.username.strip().lower()
    user = next((u for u in load_users() if u["username"] == username), None)
    if user is None or not _verify_password(body.password, user["passwordHash"]):
        raise HTTPException(401, "Invalid username or password")
    token = secrets.token_urlsafe(32)
    TOKENS[token] = username
    return {"token": token, "username": username, "role": user["role"]}


@app.get("/api/me")
def current_session(user: dict[str, str] = Depends(auth)):
    return {"username": user["username"], "role": user["role"]}


@app.post("/api/classify", dependencies=[Depends(auth)])
def classify_diagnoses(body: ClassifyRequest):
    return classify_text(body.text)


@app.post("/api/logout", dependencies=[Depends(auth)])
def logout(authorization: str | None = Header(default=None)):
    TOKENS.pop(authorization.removeprefix("Bearer "), None)
    return {"ok": True}


@app.get("/api/users", dependencies=[Depends(admin_auth)])
def users():
    return [{"username": u["username"], "role": u["role"]} for u in load_users()]


@app.post("/api/users", dependencies=[Depends(admin_auth)])
def create_user(body: NewUser):
    username = body.username.strip().lower()
    if not re.fullmatch(r"[a-z0-9._-]{3,32}", username):
        raise HTTPException(422, "Username must be 3–32 characters: letters, numbers, dot, underscore, or hyphen")
    if len(body.password) < 8:
        raise HTTPException(422, "Password must contain at least 8 characters")
    if body.role not in {"admin", "user"}:
        raise HTTPException(422, "Role must be admin or user")
    rows = load_users()
    if any(u["username"] == username for u in rows):
        raise HTTPException(409, "Username already exists")
    rows.append({"username": username, "role": body.role, "passwordHash": _password_hash(body.password)})
    save_users(rows)
    return {"username": username, "role": body.role}


@app.delete("/api/users/{username}")
def delete_user(username: str, current: dict[str, str] = Depends(admin_auth)):
    username = username.lower()
    rows = load_users()
    target = next((u for u in rows if u["username"] == username), None)
    if target is None:
        raise HTTPException(404, "Account not found")
    if username == current["username"]:
        raise HTTPException(400, "You cannot remove your own account")
    if target["role"] == "admin" and sum(u["role"] == "admin" for u in rows) <= 1:
        raise HTTPException(400, "At least one administrator account must remain")
    save_users([u for u in rows if u["username"] != username])
    for token, owner in list(TOKENS.items()):
        if owner == username:
            TOKENS.pop(token, None)
    return {"ok": True}


@app.post("/api/users/{username}/reset-password")
def reset_user_password(username: str, body: PasswordReset, current: dict[str, str] = Depends(admin_auth)):
    if len(body.password) < 8:
        raise HTTPException(422, "Password must contain at least 8 characters")
    rows = load_users()
    target = next((u for u in rows if u["username"] == username.lower()), None)
    if target is None:
        raise HTTPException(404, "Account not found")
    target["passwordHash"] = _password_hash(body.password)
    save_users(rows)
    for token, owner in list(TOKENS.items()):
        if owner == target["username"]:
            TOKENS.pop(token, None)
    return {"ok": True}


@app.get("/api/records", dependencies=[Depends(auth)])
def records():
    return read_records()


@app.post("/api/records", dependencies=[Depends(auth)])
def save_record(body: Record):
    rows = read_records()
    row = body.model_dump()
    row["id"] = body.id or secrets.token_hex(8)
    old = next((i for i, x in enumerate(rows) if x["id"] == row["id"]), None)
    if old is None:
        rows.insert(0, row)
    else:
        rows[old] = row
    write_records(rows)
    return row


@app.post("/api/records/bulk", dependencies=[Depends(auth)])
def save_records_bulk(records: list[Record]):
    if len(records) > 5000:
        raise HTTPException(413, "Too many records in one save")
    rows = read_records()
    known_cases = {row.get("caseId") for row in rows if row.get("caseId")}
    saved, skipped = [], 0
    for item in records:
        row = item.model_dump()
        if row.get("caseId") and row["caseId"] in known_cases:
            skipped += 1
            continue
        row["id"] = row.get("id") or secrets.token_hex(8)
        rows.insert(0, row)
        if row.get("caseId"):
            known_cases.add(row["caseId"])
        saved.append(row)
    if saved:
        write_records(rows)
    return {"saved": len(saved), "skipped": skipped, "records": saved}


@app.post("/api/import-cases", dependencies=[Depends(auth)])
async def import_cases(file: UploadFile = File(...)):
    name = (file.filename or "").lower()
    if not name.endswith(".xlsx"):
        raise HTTPException(415, "Choose an .xlsx workbook")
    raw = await file.read()
    if len(raw) > 30 * 1024 * 1024:
        raise HTTPException(413, "Workbook exceeds 30 MB limit")
    try:
        from openpyxl import load_workbook
        workbook = load_workbook(io.BytesIO(raw), data_only=True, read_only=True)
    except Exception as exc:
        raise HTTPException(422, "Could not read the Excel workbook") from exc
    from app.classifier import classify_text
    existing = read_records()
    known_case_ids = {row.get("caseId") for row in existing if row.get("caseId")}
    created, skipped, suggested_count, review_count, found_case_sheet = [], 0, 0, 0, False
    aliases = {
        "case": {"case id", "case no", "case number", "caseid", "case"},
        "patient": {"patient", "patient name", "name", "姓名", "病患姓名"},
        "chart": {"chart no", "chart number", "medical record no", "medical record number", "病歷號碼", "病歷編號"},
        "diagnosis": {"diagnosis", "diagnosis text", "diagnosis summary", "clinical diagnosis", "診斷", "診斷內容"},
        "date": {"exam date", "examination date", "date", "檢查日期"},
        "physician": {"physician", "doctor", "醫師"},
    }
    def norm_header(value: Any) -> str:
        return re.sub(r"[^a-z0-9\u4e00-\u9fff]+", " ", str(value or "").strip().casefold()).strip()
    for sheet in workbook.worksheets:
        iterator = sheet.iter_rows(values_only=True)
        try:
            raw_headers = list(next(iterator))
        except StopIteration:
            continue
        headers = [norm_header(value) for value in raw_headers]
        columns = {key: next((i for i, header in enumerate(headers) if header in values), None) for key, values in aliases.items()}
        if columns["case"] is None:
            continue
        found_case_sheet = True
        # Excel often represents a patient with several cases using vertically
        # merged cells. openpyxl exposes only the first cell's value, so retain
        # the patient identifiers for subsequent Case ID rows in that block.
        carried = {"patient": "", "chart": ""}
        for values in iterator:
            raw_case = values[columns["case"]] if columns["case"] < len(values) else None
            if raw_case is None or not str(raw_case).strip() or str(raw_case).strip().casefold() == "case id":
                continue
            raw_case_value = str(int(raw_case)) if isinstance(raw_case, float) and raw_case.is_integer() else str(raw_case).strip()
            # Some workbooks put several case identifiers in one cell.
            case_values = [part.strip() for part in re.split(r"[\r\n,，;；|]+", raw_case_value) if part.strip()]
            cell = lambda key: values[columns[key]] if columns[key] is not None and columns[key] < len(values) else None
            for key in carried:
                current = str(cell(key) or "").strip()
                if current:
                    carried[key] = current
            patient = carried["patient"]
            diagnosis = str(cell("diagnosis") or "").strip()
            chart = carried["chart"]
            exam_date = cell("date")
            if isinstance(exam_date, datetime):
                exam_date = exam_date.strftime("%Y-%m-%d")
            else:
                exam_date = str(exam_date or "").strip()
            physician = str(cell("physician") or "").strip()
            if not diagnosis:
                excluded_cols = {columns[key] for key in ("case", "patient", "chart", "date", "physician") if columns[key] is not None}
                diagnosis = "; ".join(f"{raw_headers[i]}: {value}" for i, value in enumerate(values) if i not in excluded_cols and i < len(raw_headers) and value is not None and str(value).strip())
            source_text = " | ".join(f"{raw_headers[i]}: {value}" for i, value in enumerate(values) if i < len(raw_headers) and value is not None and str(value).strip())
            result = classify_text(diagnosis) if diagnosis else {"suggestions": [], "review": []}
            codes = [suggestion["value"] for suggestion in result["suggestions"]]
            labels = [suggestion["label"] for suggestion in result["suggestions"]]
            for case_value in case_values:
                case_id = f"{sheet.title}-{case_value}"
                if case_id in known_case_ids:
                    skipped += 1
                    continue
                record = {
                    "id": secrets.token_hex(8), "caseId": case_id, "patientName": patient, "chartNo": chart,
                    "examDate": exam_date, "physician": physician, "functional": [], "diagnoses": codes,
                    "diagnosisLabels": labels, "classificationReview": result.get("review", []),
                    "classificationStatus": "suggested" if codes else "review", "diagnosisNotes": diagnosis, "management": [],
                    "medication": "", "sourceText": source_text,
                }
                known_case_ids.add(case_id)
                created.append(record)
                if codes:
                    suggested_count += 1
                else:
                    review_count += 1
    workbook.close()
    if not found_case_sheet:
        raise HTTPException(422, "No worksheet with a CASE ID column was found")
    return {"imported": len(created), "skipped": skipped, "suggested": suggested_count, "needsReview": review_count, "records": created}


@app.delete("/api/records/{record_id}", dependencies=[Depends(auth)])
def delete_record(record_id: str):
    rows = read_records()
    kept = [row for row in rows if row.get("id") != record_id]
    if len(kept) == len(rows):
        raise HTTPException(404, "Record not found")
    write_records(kept)
    return {"ok": True}


@app.post("/api/extract", dependencies=[Depends(auth)])
async def extract(file: UploadFile = File(...)):
    name = (file.filename or "").lower()
    raw = await file.read()
    if len(raw) > 15 * 1024 * 1024:
        raise HTTPException(413, "File exceeds 15 MB limit")
    text = ""
    try:
        if name.endswith(".xlsx"):
            from openpyxl import load_workbook
            wb = load_workbook(io.BytesIO(raw), data_only=True, read_only=True)
            text = "\n".join(" | ".join(str(c) for c in row if c is not None) for ws in wb for row in ws.iter_rows(values_only=True) if any(c is not None for c in row))
        elif name.endswith(".xls"):
            try:
                import xlrd
                book = xlrd.open_workbook(file_contents=raw)
                text = "\n".join(" | ".join(str(c) for c in row if c not in (None, "")) for sheet in book.sheets() for row in sheet.get_rows() if any(c not in (None, "") for c in row))
            except ImportError as exc:
                raise HTTPException(415, "Legacy .xls support is not installed. Save as .xlsx and upload again.") from exc
        elif name.endswith(".docx"):
            from docx import Document
            doc = Document(io.BytesIO(raw))
            text = "\n".join([p.text for p in doc.paragraphs] + [" | ".join(c.text for c in row.cells) for t in doc.tables for row in t.rows])
        elif name.endswith(".pdf"):
            from pypdf import PdfReader
            text = "\n".join(p.extract_text() or "" for p in PdfReader(io.BytesIO(raw)).pages)
            if not text.strip():
                try:
                    import pytesseract
                    from pdf2image import convert_from_bytes
                    pages = convert_from_bytes(raw, dpi=180, first_page=1, last_page=20)
                    text = "\n".join(pytesseract.image_to_string(page, lang="eng+chi_tra") for page in pages)
                except Exception as exc:
                    raise HTTPException(422, "Scanned PDF OCR needs Tesseract, eng/chi_tra language data, and Poppler.") from exc
        elif name.endswith((".png", ".jpg", ".jpeg", ".tif", ".tiff")):
            try:
                import pytesseract
                from PIL import Image
                text = pytesseract.image_to_string(Image.open(io.BytesIO(raw)), lang="eng+chi_tra")
            except Exception as exc:
                raise HTTPException(422, "Image OCR unavailable. Install Tesseract with eng and chi_tra language data.") from exc
        elif name.endswith(".doc"):
            raise HTTPException(415, "Legacy .doc is unsupported; save as .docx and upload again")
        else:
            raise HTTPException(415, "Supported: xlsx, xls, docx, pdf, png, jpg")
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(422, f"Could not extract document: {exc}") from exc
    clipped = text[:50000]
    return {"filename": file.filename, "text": clipped, "fields": extract_basic_fields(clipped)}


def selected_records(ids: list[str]) -> list[dict[str, Any]]:
    rows = read_records()
    if not ids:
        return rows
    selected = set(ids)
    return [row for row in rows if str(row.get("id")) in selected]


def record_export_rows(rows: list[dict[str, Any]]) -> list[list[str]]:
    header = ["Record ID", "Patient Name", "Chart No", "Exam Date", "Physician", "Functional Classifications", "Diagnostic Items", "Treatment Plan", "Medication"]
    return [header] + [[str(r.get("id", "")), str(r.get("patientName", "")), str(r.get("chartNo", "")), str(r.get("examDate", "")), str(r.get("physician", "")), "; ".join(r.get("functional", [])), "; ".join(r.get("diagnosisLabels", []) + ([r["diagnosisNotes"]] if r.get("diagnosisNotes") else [])), "; ".join(r.get("management", [])), str(r.get("medication", ""))] for r in rows]


@app.get("/api/export", dependencies=[Depends(auth)])
def export(format: str = "xlsx", ids: list[str] = Query(default=[])):
    rows = selected_records(ids)
    if ids and not rows:
        raise HTTPException(404, "No matching records selected")
    if format == "docx":
        from docx import Document
        document = Document()
        document.add_heading("Hospital Rehabilitation Checklist Records", 0)
        for record in rows:
            document.add_heading(record.get("patientName") or "Unnamed patient", level=1)
            table = document.add_table(rows=0, cols=2)
            table.style = "Light Shading Accent 1"
            fields = [("Record ID", record.get("id", "")), ("Case ID", record.get("caseId", "")), ("Chart No", record.get("chartNo", "")), ("Exam Date", record.get("examDate", "")), ("Physician", record.get("physician", "")), ("Functional classifications", "; ".join(record.get("functional", []))), ("Diagnosis", "; ".join(record.get("diagnosisLabels", []) + ([record["diagnosisNotes"]] if record.get("diagnosisNotes") else []))), ("Treatment plan", "; ".join(record.get("management", []))), ("Medication / notes", record.get("medication", ""))]
            for label, value in fields:
                cells = table.add_row().cells
                cells[0].text, cells[1].text = label, str(value)
        buf = io.BytesIO(); document.save(buf); buf.seek(0)
        return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document", headers={"Content-Disposition": 'attachment; filename="hospital_checklist.docx"'})
    if format == "pdf":
        try:
            from reportlab.lib import colors
            from reportlab.lib.pagesizes import landscape, letter
            from reportlab.lib.styles import getSampleStyleSheet
            from reportlab.lib.units import inch
            from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
            from reportlab.pdfbase import pdfmetrics
            from reportlab.pdfbase.ttfonts import TTFont
        except ImportError as exc:
            raise HTTPException(503, "PDF export dependency is missing; install reportlab") from exc
        buf = io.BytesIO()
        doc = SimpleDocTemplate(buf, pagesize=landscape(letter), rightMargin=0.45*inch, leftMargin=0.45*inch, topMargin=0.45*inch, bottomMargin=0.45*inch)
        styles = getSampleStyleSheet()
        cjk_font = Path("/System/Library/Fonts/PingFang.ttc")
        if cjk_font.exists():
            pdfmetrics.registerFont(TTFont("ChecklistCJK", str(cjk_font), subfontIndex=0))
            for style in styles.byName.values():
                style.fontName = "ChecklistCJK"
        story = [Paragraph("Hospital Rehabilitation Checklist Records", styles["Title"]), Spacer(1, 12)]
        for record in rows:
            story.append(Paragraph(str(record.get("patientName") or "Unnamed patient"), styles["Heading2"]))
            data = [["Field", "Value"]] + [[label, str(value or "")] for label, value in [("Record ID", record.get("id")), ("Case ID", record.get("caseId")), ("Chart No", record.get("chartNo")), ("Exam Date", record.get("examDate")), ("Physician", record.get("physician")), ("Functional classifications", "; ".join(record.get("functional", []))), ("Diagnosis", "; ".join(record.get("diagnosisLabels", []) + ([record["diagnosisNotes"]] if record.get("diagnosisNotes") else []))), ("Treatment plan", "; ".join(record.get("management", []))), ("Medication / notes", record.get("medication"))]]
            table = Table([[Paragraph(escape(str(cell)), styles["BodyText"]) for cell in row] for row in data], colWidths=[1.8*inch, 8.0*inch], repeatRows=1)
            table.setStyle(TableStyle([("BACKGROUND", (0,0),(-1,0), colors.HexColor("#bd7137")), ("TEXTCOLOR", (0,0),(-1,0), colors.white), ("GRID", (0,0),(-1,-1), .4, colors.HexColor("#dedede")), ("VALIGN", (0,0),(-1,-1), "TOP"), ("PADDING", (0,0),(-1,-1), 6)]))
            story.extend([table, Spacer(1, 12)])
        doc.build(story); buf.seek(0)
        return StreamingResponse(buf, media_type="application/pdf", headers={"Content-Disposition": 'attachment; filename="hospital_checklist.pdf"'})
    if format != "xlsx":
        raise HTTPException(422, "Format must be xlsx, docx, or pdf")
    wb = Workbook()
    ws = wb.active
    ws.title = "Checklist records"
    headers = ["Record ID", "Case ID", "Patient Name", "Chart No", "Exam Date", "Physician", "Functional Classifications", "Diagnostic Items", "Treatment Plan", "Medication"]
    headers.extend(f"Diagnosis {code} — {name}" for code, name in DIAGNOSIS_GROUPS.items())
    ws.append(headers)
    for r in rows:
        diagnoses = r.get("diagnoses", [])
        diagnosis_items = r.get("diagnosisLabels", []) + ([r["diagnosisNotes"]] if r.get("diagnosisNotes") else [])
        flags = [int(any(item.startswith(f"diag-{group_id}-") for item in diagnoses)) for group_id in DIAGNOSIS_GROUPS]
        ws.append([r.get("id"), r.get("caseId"), r.get("patientName"), r.get("chartNo"), r.get("examDate"), r.get("physician"), "; ".join(r.get("functional", [])), "; ".join(diagnosis_items or diagnoses), "; ".join(r.get("management", [])), r.get("medication", ""), *flags])
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="BD7137")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    widths = [18, 18, 24, 18, 16, 22, 40, 55, 38, 30] + [24] * len(DIAGNOSIS_GROUPS)
    for index, width in enumerate(widths, start=1):
        ws.column_dimensions[ws.cell(1, index).column_letter].width = width
    buf = io.BytesIO(); wb.save(buf); buf.seek(0)
    return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers={"Content-Disposition": 'attachment; filename="hospital_checklist.xlsx"'})
